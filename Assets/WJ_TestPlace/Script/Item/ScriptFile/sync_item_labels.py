# -*- coding: utf-8 -*-
"""
아이템 데이터 시트(ItemDataTable.xlsx)와 번역 라벨 시트(ItemDataLabel.xlsx)의 itemId를 동일하게 맞추고,
KOR/ENG 시트의 행 순서를 ItemDataTable.xlsx에 나온 순서(무기 -> 방어구 -> 포션 -> 유물)로 재배열한다.

동작:
  1. ItemDataTable.xlsx의 WeaponDefinitions/ArmorDefinitions/PotionDefinitions/RelicDefinitions 시트를
     이 순서로 읽어서 "기준 itemId 순서"와 각 itemId의 한글 itemName/description을 모은다.
  2. ItemDataLabel.xlsx의 KOR/ENG 시트를 읽어서 기존 번역 내용을 itemId 기준으로 보존한다.
  3. 기준 목록에 없는데 라벨 시트에만 있는 itemId(더 이상 존재하지 않는 아이템)는 제거한다.
  4. 기준 목록에는 있는데 라벨 시트에 없는 itemId(번역이 아직 없는 새 아이템)는 새로 추가한다.
     - KOR 시트: ItemDataTable.xlsx의 한글 itemName/description을 그대로 채워 넣는다.
     - ENG 시트: itemId만 채우고 itemName/description은 비워둔다 (번역 필요 표시).
  5. 두 시트 모두 기준 순서대로 다시 정렬해서 저장한다.

실행 전 두 파일이 Excel에서 열려있지 않은지(~$ 잠금 파일) 확인한다.

사용법: python sync_item_labels.py
"""
import os
import sys

import openpyxl
from openpyxl.styles import Font

PROJECT_ROOT = os.path.dirname(os.path.abspath(__file__))
MAIN_TABLE_PATH = os.path.join(
    PROJECT_ROOT, "Assets", "Resources", "DataFiles", "ItemData",
    "1. ExcelFile", "ItemDataTable.xlsx")
LABEL_PATH = os.path.join(
    PROJECT_ROOT, "Assets", "Resources", "DataFiles", "ItemData",
    "1. ExcelFile", "ItemDataLabel.xlsx")

# ItemDataTable.xlsx를 읽는 순서. RelicDefinitions는 아직 비어있을 수 있지만 그대로 둔다.
MAIN_SHEET_ORDER = [
    "WeaponDefinitions",
    "ArmorDefinitions",
    "PotionDefinitions",
    "RelicDefinitions",
]

TYPE_HINT_TOKENS = {"string", "int", "float", "bool"}


def is_lock_present(path):
    directory = os.path.dirname(path)
    lock_name = "~$" + os.path.basename(path)
    return os.path.exists(os.path.join(directory, lock_name))


def read_main_table(path):
    """기준 itemId 순서 목록과 {itemId: (itemName, description)} 딕셔너리를 반환한다."""
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)

    order = []
    kor_content = {}

    for sheet_name in MAIN_SHEET_ORDER:
        if sheet_name not in wb.sheetnames:
            continue

        ws = wb[sheet_name]
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue

        header = [str(c).strip() if c is not None else "" for c in rows[0]]
        if "itemId" not in header:
            print(f"  경고: {sheet_name} 시트에 itemId 컬럼이 없어 건너뜁니다.")
            continue

        id_idx = header.index("itemId")
        name_idx = header.index("itemName") if "itemName" in header else None
        desc_idx = header.index("description") if "description" in header else None

        for row in rows[1:]:
            if row is None or len(row) <= id_idx:
                continue

            # 1행 뒤에 타입 힌트 행이 올 수도 있으니 걸러낸다.
            non_empty = [str(v).strip() for v in row if v is not None and str(v).strip() != ""]
            if non_empty and all(v.lower() in TYPE_HINT_TOKENS for v in non_empty):
                continue

            item_id = row[id_idx]
            if item_id is None or str(item_id).strip() == "":
                continue
            item_id = str(item_id).strip()

            item_name = row[name_idx] if name_idx is not None and name_idx < len(row) else ""
            description = row[desc_idx] if desc_idx is not None and desc_idx < len(row) else ""

            order.append(item_id)
            kor_content[item_id] = (item_name or "", description or "")

    return order, kor_content


def read_label_sheet(ws):
    """{itemId: (itemName, description)} 딕셔너리로 반환한다. 헤더 행은 제외."""
    rows = list(ws.iter_rows(values_only=True))
    content = {}
    if not rows:
        return content

    header = [str(c).strip() if c is not None else "" for c in rows[0]]
    if "itemId" not in header:
        return content

    id_idx = header.index("itemId")
    name_idx = header.index("itemName") if "itemName" in header else None
    desc_idx = header.index("description") if "description" in header else None

    for row in rows[1:]:
        if row is None or len(row) <= id_idx:
            continue
        item_id = row[id_idx]
        if item_id is None or str(item_id).strip() == "":
            continue
        item_id = str(item_id).strip()

        item_name = row[name_idx] if name_idx is not None and name_idx < len(row) else ""
        description = row[desc_idx] if desc_idx is not None and desc_idx < len(row) else ""
        content[item_id] = (item_name or "", description or "")

    return content


def write_label_sheet(ws, ordered_rows):
    """헤더(1행)만 남기고 나머지를 지운 뒤, ordered_rows(itemId, itemName, description) 목록으로 다시 채운다."""
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)

    for item_id, item_name, description in ordered_rows:
        ws.append([item_id, item_name, description])

    header_font = Font(bold=True)
    for cell in ws[1]:
        cell.font = header_font

    ws.column_dimensions["A"].width = 32
    ws.column_dimensions["B"].width = 22
    ws.column_dimensions["C"].width = 40


def main():
    if not os.path.exists(MAIN_TABLE_PATH):
        print("ItemDataTable.xlsx를 찾을 수 없습니다:", MAIN_TABLE_PATH)
        sys.exit(1)
    if not os.path.exists(LABEL_PATH):
        print("ItemDataLabel.xlsx를 찾을 수 없습니다:", LABEL_PATH)
        sys.exit(1)

    if is_lock_present(MAIN_TABLE_PATH):
        print("ItemDataTable.xlsx가 Excel에서 열려있는 것 같습니다. 닫고 다시 실행해주세요.")
        sys.exit(1)
    if is_lock_present(LABEL_PATH):
        print("ItemDataLabel.xlsx가 Excel에서 열려있는 것 같습니다. 닫고 다시 실행해주세요.")
        sys.exit(1)

    master_order, kor_source = read_main_table(MAIN_TABLE_PATH)
    if not master_order:
        print("ItemDataTable.xlsx에서 유효한 itemId를 하나도 찾지 못했습니다. 중단합니다.")
        sys.exit(1)

    master_set = set(master_order)

    wb = openpyxl.load_workbook(LABEL_PATH)
    if "KOR" not in wb.sheetnames or "ENG" not in wb.sheetnames:
        print("ItemDataLabel.xlsx에 KOR/ENG 시트가 모두 있어야 합니다.")
        sys.exit(1)

    kor_existing = read_label_sheet(wb["KOR"])
    eng_existing = read_label_sheet(wb["ENG"])

    added = [item_id for item_id in master_order if item_id not in kor_existing and item_id not in eng_existing]
    removed_kor = [item_id for item_id in kor_existing if item_id not in master_set]
    removed_eng = [item_id for item_id in eng_existing if item_id not in master_set]
    removed = sorted(set(removed_kor) | set(removed_eng))

    kor_rows = []
    eng_rows = []
    for item_id in master_order:
        if item_id in kor_existing:
            name, desc = kor_existing[item_id]
        else:
            name, desc = kor_source.get(item_id, ("", ""))
        kor_rows.append((item_id, name, desc))

        if item_id in eng_existing:
            name, desc = eng_existing[item_id]
        else:
            name, desc = "", ""
        eng_rows.append((item_id, name, desc))

    write_label_sheet(wb["KOR"], kor_rows)
    write_label_sheet(wb["ENG"], eng_rows)
    wb.save(LABEL_PATH)

    print("===== ItemDataLabel.xlsx 동기화 완료 =====")
    print(f"기준(ItemDataTable.xlsx) itemId 개수: {len(master_order)}")
    print(f"순서: {' -> '.join(master_order)}")
    if added:
        print(f"새로 추가됨 ({len(added)}개): {', '.join(added)}")
        print("  (KOR은 ItemDataTable의 한글 itemName/description으로 채웠고, ENG는 비워뒀습니다. 번역이 필요합니다.)")
    else:
        print("새로 추가된 항목 없음")
    if removed:
        print(f"제거됨 ({len(removed)}개, ItemDataTable에 더 이상 없음): {', '.join(removed)}")
    else:
        print("제거된 항목 없음")


if __name__ == "__main__":
    main()
